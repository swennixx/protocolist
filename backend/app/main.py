import asyncio
import io
import json
import os
import re
import shutil
import threading
import uuid
from contextlib import asynccontextmanager
from datetime import date

from docx import Document
from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, Response, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select
from starlette.concurrency import run_in_threadpool, iterate_in_threadpool

from . import llm, ml
from .db import DEFAULT_SETTINGS, Meeting, Segment, Session, Setting, Speaker, Task, get_settings
from .pipeline import STAGES, meeting_dir, paths
from .search import search as run_search
from .text import fmt_time, to_srt, to_vtt

MAX_SIZE = 2 * 1024**3
EXTENSIONS = {".mp3", ".m4a", ".wav", ".ogg", ".oga", ".opus", ".webm", ".mp4", ".mov", ".mkv", ".flac", ".aac"}

@asynccontextmanager
async def lifespan(_app):
    # Load the text-embedding model in the background, so the first search does not wait for it.
    threading.Thread(target=lambda: ml.embed(["прогрев"], query=True), daemon=True).start()
    yield


app = FastAPI(title="Протоколист API", lifespan=lifespan)


def get_db():
    with Session() as db:
        yield db


def load(db, mid: str) -> Meeting:
    m = db.get(Meeting, mid) if re.fullmatch(r"[0-9a-f]{32}", mid) else None
    if not m:
        raise HTTPException(404, "Встреча не найдена")
    return m


def speaker_json(s: Speaker) -> dict:
    return {"id": s.id, "name": s.name, "color": s.idx % 5}


def brief(m: Meeting) -> dict:
    return {
        "id": m.id, "title": m.title, "date": m.created_at.isoformat(), "status": m.status, "stage": m.stage,
        "stage_progress": m.stage_progress, "error": m.error, "duration": m.duration,
        "size": m.size, "received": m.received,
        "speakers": [speaker_json(s) for s in m.speakers],
        "summary": m.summary[:1], "decisions": len(m.decisions), "open_tasks": sum(not t.done for t in m.tasks),
    }


def full(m: Meeting) -> dict:
    return {
        **brief(m),
        "language": m.language, "has_audio": m.has_audio, "peaks": m.peaks, "summary": m.summary, "topics": m.topics,
        "timings": m.timings,
        "segments": [
            {"id": s.id, "speaker": s.speaker_id, "start": s.start, "end": s.end, "text": s.text,
             "clean": s.text_clean, "edited": s.original_clean is not None}
            for s in m.segments
        ],
        "decisions": [{"text": d.text, "seg": d.seg} for d in m.decisions],
        "tasks": [
            {"id": t.id, "text": t.text, "owner": t.owner_id, "owner_name": t.owner_name,
             "due": t.due.isoformat() if t.due else None, "seg": t.seg, "done": t.done}
            for t in m.tasks
        ],
    }


# --- meetings ---

class NewMeeting(BaseModel):
    title: str = Field(min_length=1, max_length=300)
    filename: str = Field(max_length=300)
    size: int = Field(gt=0, le=MAX_SIZE)
    language: str | None = Field(default=None, pattern="^(ru|en)$")
    num_speakers: int | None = Field(default=None, ge=1, le=20)


@app.get("/api/meetings")
def list_meetings(db=Depends(get_db)):
    return [brief(m) for m in db.scalars(select(Meeting).order_by(Meeting.created_at.desc()))]


@app.post("/api/meetings", status_code=201)
def create_meeting(body: NewMeeting, db=Depends(get_db)):
    if os.path.splitext(body.filename.lower())[1] not in EXTENSIONS:
        raise HTTPException(400, "Неподдерживаемый формат файла")
    m = Meeting(id=uuid.uuid4().hex, title=body.title.strip(), filename=body.filename, size=body.size,
                language=body.language, num_speakers=body.num_speakers)
    db.add(m)
    db.commit()
    os.makedirs(meeting_dir(m.id), exist_ok=True)
    return brief(m)


@app.put("/api/meetings/{mid}/upload")
async def upload_chunk(mid: str, request: Request, offset: int = Query(ge=0)):
    """Resumable upload: the client sends chunks in order; on a mismatch it resumes from `received`."""
    with Session() as db:
        m = load(db, mid)
        if m.status != "uploading":
            raise HTTPException(409, "Файл уже загружен")
        if offset != m.received:
            raise HTTPException(409, detail={"received": m.received})
    path = paths(mid)["upload"]
    written = 0
    with open(path, "r+b" if os.path.exists(path) else "wb") as f:
        f.seek(offset)
        f.truncate()
        async for chunk in request.stream():
            written += len(chunk)
            if offset + written > m.size:
                raise HTTPException(413, "Файл больше заявленного размера")
            f.write(chunk)
    with Session.begin() as db:
        m = load(db, mid)
        m.received = offset + written
        if m.received == m.size:
            m.status = "queued"
        return {"received": m.received, "status": m.status}


@app.get("/api/meetings/{mid}")
def get_meeting(mid: str, db=Depends(get_db)):
    return full(load(db, mid))


class MeetingPatch(BaseModel):
    title: str = Field(min_length=1, max_length=300)


@app.patch("/api/meetings/{mid}")
def rename_meeting(mid: str, body: MeetingPatch, db=Depends(get_db)):
    m = load(db, mid)
    m.title = body.title.strip()
    db.commit()
    return brief(m)


@app.delete("/api/meetings/{mid}", status_code=204)
def delete_meeting(mid: str, db=Depends(get_db)):
    m = load(db, mid)
    if m.status == "processing":
        raise HTTPException(409, "Дождитесь окончания обработки")
    db.delete(m)
    db.commit()
    shutil.rmtree(meeting_dir(mid), ignore_errors=True)


@app.post("/api/meetings/{mid}/retry")
def retry(mid: str, db=Depends(get_db)):
    m = load(db, mid)
    if m.status != "error":
        raise HTTPException(409, "Повторить можно только упавшую обработку")
    m.status, m.error = "queued", None
    db.commit()
    return brief(m)


@app.post("/api/meetings/{mid}/reprocess")
def reprocess(mid: str, stage: str = Query("asr", pattern="^(asr|diarize|analyze)$"), db=Depends(get_db)):
    """Run the pipeline again from a stage, e.g. after changing the model or the glossary.
    The source file is gone after preparation, so the earliest stage is recognition."""
    m = load(db, mid)
    if m.status not in ("done", "error"):
        raise HTTPException(409, "Встреча ещё обрабатывается")
    if not os.path.exists(paths(mid)["wav"]):
        raise HTTPException(409, "Аудио удалено после обработки — распознать заново нельзя")
    m.stage, m.stage_progress, m.status, m.error = STAGES.index(stage), 0, "queued", None
    db.commit()
    return brief(m)


@app.get("/api/meetings/{mid}/events")
async def events(mid: str, request: Request):
    """SSE stream of processing progress. ponytail: polls the DB twice a second; LISTEN/NOTIFY if many clients watch."""

    def snapshot():
        with Session() as db:
            m = load(db, mid)
            return {"status": m.status, "stage": m.stage, "stage_progress": m.stage_progress, "error": m.error,
                    "received": m.received, "size": m.size}

    await run_in_threadpool(snapshot)  # 404 before the stream starts

    async def gen():
        last = None
        while not await request.is_disconnected():
            snap = await run_in_threadpool(snapshot)
            if snap != last:
                yield f"data: {json.dumps(snap)}\n\n"
                last = snap
            if snap["status"] in ("done", "error"):
                return
            await asyncio.sleep(0.5)

    return StreamingResponse(gen(), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


@app.get("/api/meetings/{mid}/audio")
def audio(mid: str, db=Depends(get_db)):
    load(db, mid)
    path = paths(mid)["play"]
    if not os.path.exists(path):
        raise HTTPException(404, "Аудио недоступно")
    return FileResponse(path, media_type="audio/mp4")


# --- edits ---

class Name(BaseModel):
    name: str = Field(min_length=1, max_length=100)


@app.patch("/api/speakers/{sid}")
def rename_speaker(sid: int, body: Name, db=Depends(get_db)):
    s = db.get(Speaker, sid) or _404()
    s.name = body.name.strip()
    db.commit()
    return speaker_json(s)


class SegmentPatch(BaseModel):
    clean: str = Field(min_length=1, max_length=20000)


@app.patch("/api/segments/{sid}")
def edit_segment(sid: int, body: SegmentPatch, db=Depends(get_db)):
    s = db.get(Segment, sid) or _404()
    if s.original_clean is None:
        s.original_clean = s.text_clean
    s.text_clean = body.clean.strip()
    db.commit()
    return {"id": s.id, "clean": s.text_clean, "edited": True}


class TaskPatch(BaseModel):
    done: bool


@app.patch("/api/tasks/{tid}")
def edit_task(tid: int, body: TaskPatch, db=Depends(get_db)):
    t = db.get(Task, tid) or _404()
    t.done = body.done
    db.commit()
    return {"id": t.id, "done": t.done}


def _404():
    raise HTTPException(404, "Не найдено")


# --- search & ask ---

@app.get("/api/search")
def search(q: str = Query(max_length=500), mode: str = Query("smart", pattern="^(smart|exact)$"), db=Depends(get_db)):
    return run_search(db, q, mode)


class Question(BaseModel):
    question: str = Field(min_length=2, max_length=1000)
    meeting_id: str | None = None


@app.post("/api/ask")
def ask(body: Question):
    """RAG: hybrid retrieval, then a streamed answer with [n] citations. SSE events: sources, token, done."""
    with Session() as db:
        hits = run_search(db, body.question, "smart", body.meeting_id, limit=8)
        model = get_settings(db)["llm"]
        fragments = [_with_context(db, h) for h in hits]

    def gen():
        yield f"event: sources\ndata: {json.dumps(hits, ensure_ascii=False)}\n\n"
        if not hits:
            yield f"event: token\ndata: {json.dumps('В записанных встречах нет информации по этому вопросу.', ensure_ascii=False)}\n\n"
        elif model == "off":
            yield f"event: token\ndata: {json.dumps('Языковая модель выключена в настройках — вот найденные фрагменты.', ensure_ascii=False)}\n\n"
        else:
            try:
                for token in llm.answer(model, body.question, fragments):
                    yield f"event: token\ndata: {json.dumps(token, ensure_ascii=False)}\n\n"
            except Exception:
                yield f"event: token\ndata: {json.dumps(' [Языковая модель недоступна: проверьте, что запущен Ollama.]', ensure_ascii=False)}\n\n"
        yield "event: done\ndata: {}\n\n"

    return StreamingResponse(iterate_in_threadpool(gen()), media_type="text/event-stream", headers={"Cache-Control": "no-cache"})


def _with_context(db, hit: dict) -> str:
    """A found turn plus its neighbours: the reason or the reply is often one turn away."""
    rows = db.execute(
        select(Segment.idx, Segment.text_clean, Speaker.name)
        .join(Speaker, Speaker.id == Segment.speaker_id)
        .where(Segment.meeting_id == hit["meeting_id"], Segment.idx.between(hit["seg"] - 1, hit["seg"] + 1))
        .order_by(Segment.idx)
    ).all()
    lines = [f"{'→ ' if idx == hit['seg'] else ''}{name}: {text}" for idx, text, name in rows]
    return f"Встреча «{hit['meeting_title']}», {hit['date'][:10]}, {fmt_time(hit['start'])}:\n" + "\n".join(lines)


# --- export ---

@app.get("/api/meetings/{mid}/export")
def export(mid: str, format: str = Query(pattern="^(md|srt|vtt|json|docx)$"), db=Depends(get_db)):
    m = load(db, mid)
    names = {s.id: s.name for s in m.speakers}
    segs = [{"speaker": names[s.speaker_id], "start": s.start, "end": s.end, "text": s.text_clean} for s in m.segments]
    day = m.created_at.strftime("%d.%m.%Y")
    owner = lambda t: names.get(t.owner_id) or t.owner_name or "—"  # noqa: E731
    due = lambda t: f", до {t.due:%d.%m}" if t.due else ""  # noqa: E731
    fname = re.sub(r'[\\/:*?"<>|]+', "_", m.title)[:80]

    if format in ("srt", "vtt"):
        body, mime = (to_srt(segs), "application/x-subrip") if format == "srt" else (to_vtt(segs), "text/vtt")
    elif format == "json":
        body, mime = json.dumps(full(m), ensure_ascii=False, indent=2), "application/json"
    elif format == "md":
        lines = [f"# {m.title}", f"{day} · {fmt_time(m.duration)} · {', '.join(names.values())}", "## Кратко"]
        lines += [f"- {s}" for s in m.summary]
        lines += ["## Решения"] + [f"- {d.text} ({fmt_time(m.segments[d.seg].start)})" for d in m.decisions]
        lines += ["## Задачи"] + [f"- [{'x' if t.done else ' '}] {t.text} — {owner(t)}{due(t)}" for t in m.tasks]
        lines += ["## Стенограмма"] + [f"**{s['speaker']}** [{fmt_time(s['start'])}]: {s['text']}" for s in segs]
        body, mime = "\n\n".join(lines), "text/markdown"
    else:
        doc = Document()
        doc.add_heading(m.title, 0)
        doc.add_paragraph(f"{day} · {fmt_time(m.duration)} · Участники: {', '.join(names.values())}")
        for title, items in (("Кратко", m.summary), ("Решения", [d.text for d in m.decisions]),
                             ("Задачи", [f"{t.text} — {owner(t)}{due(t)}" for t in m.tasks])):
            if items:
                doc.add_heading(title, 1)
                for it in items:
                    doc.add_paragraph(it, style="List Bullet")
        doc.add_heading("Стенограмма", 1)
        for s in segs:
            p = doc.add_paragraph()
            p.add_run(f"{s['speaker']} [{fmt_time(s['start'])}]: ").bold = True
            p.add_run(s["text"])
        buf = io.BytesIO()
        doc.save(buf)
        body, mime = buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    from urllib.parse import quote
    return Response(body, media_type=mime, headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(fname)}.{format}"})


# --- settings ---

class SettingsBody(BaseModel):
    model: str = Field(pattern="^(tiny|small|medium|large-v3-turbo|large-v3)$")
    lang: str = Field(pattern="^(auto|ru|en)$")
    llm: str = Field(max_length=100)
    template: str = Field(pattern="^(short|full|protocol)$")
    glossary: list[str] = Field(max_length=200)
    deleteAudio: bool
    diarization: str = Field(default="auto", pattern="^(auto|clustering|pyannote)$")


@app.get("/api/settings")
def read_settings(db=Depends(get_db)):
    return {**get_settings(db), "pyannote_available": ml.pyannote_downloaded()}


@app.put("/api/settings")
def write_settings(body: SettingsBody, db=Depends(get_db)):
    row = db.get(Setting, 1) or Setting(id=1)
    row.data = body.model_dump()
    db.merge(row)
    db.commit()
    return read_settings(db)


@app.get("/api/health")
def health():
    return {"ok": True, "stages": STAGES, "defaults": DEFAULT_SETTINGS, "today": date.today().isoformat()}
